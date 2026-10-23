# -*- coding: utf-8 -*-
"""Bộ máy chuẩn hóa văn bản Word (.docx) – TEXO, bản 2.

Nguyên tắc:
  • Chỉ chỉnh THỂ THỨC (khổ giấy, lề, phông, cỡ chữ, giãn dòng, khoảng cách đoạn) và sửa lỗi gõ chắc chắn.
  • Giữ nguyên căn giữa / căn phải (kể cả khi căn theo Style), giữ đậm/nghiêng, không động vào số hiệu văn bản viện dẫn.
  • Đoạn có hình ảnh không đặt giãn dòng "Exactly" (tránh cắt mất ảnh).
  • Mọi thay đổi chữ được ghi vào nhật ký để người dùng soát lại.
"""
import io
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

FONT = "Times New Roman"

# --------------------------------------------------------------------------- bộ quy chuẩn
PROFILES = {
    "texo": {
        "ten": "Quy định TEXO (nội bộ)",
        "body": 13, "title": 14, "upper": 12, "table": 12,
        "line": ("exact", 17), "table_line": ("exact", 15),
        "before": 6, "after": 3, "table_before": 3, "table_after": 3,
        "cell_indent_mm": 2, "first_line_cm": 0,
        "margins": {"thuong": (20, 20, 30, 20), "letterhead": (40, 30, 30, 20)},
    },
    "nd30": {
        "ten": "Nghị định 30/2020/NĐ-CP",
        "body": 14, "title": 14, "upper": 13, "table": 13,
        "line": ("multiple", 1.15), "table_line": ("multiple", 1.0),
        "before": 6, "after": 0, "table_before": 0, "table_after": 0,
        "cell_indent_mm": 0, "first_line_cm": 1.27,
        "margins": {"thuong": (20, 20, 30, 15)},
    },
}

# Lỗi gõ chắc chắn sai – sửa an toàn
SAFE_FIXES = [
    (r"\bquy ttrình\b", "quy trình"),
    (r"\bloà việc\b", "làm việc"),
    (r"\bsai xót\b", "sai sót"),
    (r"\bko\b", "không"),
    (r"\bchình ký\b", "trình ký"),
    (r"\bthe thức\b", "thể thức"),
    (r"\btu vấn\b", "tư vấn"),
    (r"\bdu án\b", "dự án"),
    (r"\bnên kế hoạch\b", "lên kế hoạch"),
    (r"\bxử lí\b", "xử lý"),
    (r"\bquản lí\b", "quản lý"),
    (r"\bkỹ thuật\s+kỹ thuật\b", "kỹ thuật"),
]
# Dấu phẩy sau liên từ – chỉ khi đứng đầu câu, viết hoa
COMMA_RULE = re.compile(r"(^|[.;:!?]\s+)(Tuy nhiên|Do đó|Vì vậy|Bởi vậy|Theo đó|Ngoài ra|Mặt khác)\s+(?=[^,\s])")
# Ngày tháng theo NĐ 30: ngày < 10 và tháng 1, 2 thêm số 0 – chỉ trong cụm "ngày … tháng … năm …"
DATE_RULE = re.compile(r"\b(ngày)\s+(\d{1,2})\s+(tháng)\s+(\d{1,2})\s+(năm)\s+(\d{4})", re.IGNORECASE)
DOUBLE_SPACE = re.compile(r"(?<=\S) {2,}(?=\S)")
SPACE_BEFORE_PUNCT = re.compile(r"(?<=\w) +([,.;:!?])(?=\s|$)")


def _date_fix(m):
    d, mo = int(m.group(2)), int(m.group(4))
    dd = f"{d:02d}" if d < 10 else str(d)
    mm = f"{mo:02d}" if mo in (1, 2) else str(mo)
    return f"{m.group(1)} {dd} {m.group(3)} {mm} {m.group(5)} {m.group(6)}"


def fix_text(text, opts):
    new = text
    if opts.get("spell", True):
        for pat, rep in SAFE_FIXES:
            new = re.sub(pat, lambda m, r=rep: _keep_case(m.group(0), r), new, flags=re.IGNORECASE)
    if opts.get("comma", True):
        new = COMMA_RULE.sub(lambda m: f"{m.group(1)}{m.group(2)}, ", new)
    if opts.get("date", True):
        new = DATE_RULE.sub(_date_fix, new)
    if opts.get("spaces", True):
        new = DOUBLE_SPACE.sub(" ", new)
        new = SPACE_BEFORE_PUNCT.sub(r"\1", new)
    return new


def _keep_case(src, rep):
    return rep[:1].upper() + rep[1:] if src[:1].isupper() else rep


# --------------------------------------------------------------------------- tiện ích docx
def _eff_align(p):
    if p.alignment is not None:
        return p.alignment
    s = p.style
    while s is not None:
        if s.paragraph_format.alignment is not None:
            return s.paragraph_format.alignment
        s = s.base_style
    return None


def _rfonts(rpr):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), FONT)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]


def _font_run(run, size=None):
    _rfonts(run._r.get_or_add_rPr())
    if size:
        run.font.size = Pt(size)


def _has_image(p):
    return bool(p._p.xpath(".//w:drawing | .//w:pict | .//w:object"))


def _is_list(p):
    return p._p.pPr is not None and p._p.pPr.numPr is not None


def _no_contextual(p):
    pPr = p._p.get_or_add_pPr()
    el = pPr.find(qn("w:contextualSpacing"))
    if el is not None:
        pPr.remove(el)


def _line(pf, spec):
    kind, val = spec
    if kind == "exact":
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(val)
    else:
        pf.line_spacing = val


def _is_upper_line(t):
    letters = [c for c in t if c.isalpha()]
    return len(letters) >= 3 and all(c.isupper() for c in letters)


# --------------------------------------------------------------------------- chuẩn hóa
def standardize(data, profile="texo", paper="thuong", opts=None):
    """data: bytes .docx → (bytes .docx đã chuẩn hóa, nhật ký dict)."""
    opts = opts or {}
    P = PROFILES[profile]
    doc = docx.Document(io.BytesIO(data))
    log = {"bo_quy_chuan": P["ten"] + (" – giấy Letterhead" if paper == "letterhead" else ""),
           "le_cu": [], "doan": 0, "can_deu": 0, "giu_can": 0, "anh": 0, "thut_dau_dong": 0,
           "o_bang": 0, "sua_chu": []}

    # 1. Khổ giấy & lề (giữ hướng ngang nếu trang đang ngang)
    t, b, l, r = P["margins"].get(paper, P["margins"]["thuong"])
    for s in doc.sections:
        log["le_cu"].append(tuple(round(x.mm) if x is not None else None
                                  for x in (s.top_margin, s.bottom_margin, s.left_margin, s.right_margin)))
        landscape = s.page_width is not None and s.page_height is not None and s.page_width > s.page_height
        s.page_width, s.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Mm(t), Mm(b), Mm(l), Mm(r)
        for hf in (s.header, s.footer):
            if hf is not None and not hf.is_linked_to_previous:
                for p in hf.paragraphs:
                    for run in p.runs:
                        _font_run(run)

    # 2. Phông mặc định của văn bản
    try:
        st = doc.styles["Normal"]
        st.font.name = FONT
        _rfonts(st.element.get_or_add_rPr())
        st.font.size = Pt(P["body"])
    except Exception:
        pass

    def fix_runs(p, where):
        if not opts.get("text", True):
            return
        for run in p.runs:
            if run.text and run.text.strip():
                new = fix_text(run.text, opts)
                if new != run.text:
                    log["sua_chu"].append({"Vị trí": where, "Trước": run.text.strip()[:160],
                                           "Sau": new.strip()[:160]})
                    run.text = new

    # 3. Đoạn văn ngoài bảng
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        log["doan"] += 1
        fix_runs(p, f"Đoạn {i}")
        pf = p.paragraph_format
        al = _eff_align(p)
        if _has_image(p):
            log["anh"] += 1
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        else:
            _line(pf, P["line"])
        pf.space_before, pf.space_after = Pt(P["before"]), Pt(P["after"])
        _no_contextual(p)
        if text and al in (None, AL.LEFT, AL.JUSTIFY, AL.DISTRIBUTE):
            p.alignment = AL.JUSTIFY
            log["can_deu"] += 1
        elif text:
            log["giu_can"] += 1
        upper = _is_upper_line(text)
        is_title = upper and len(text) < 100 and al == AL.CENTER
        if (P["first_line_cm"] and opts.get("indent", True) and text and not upper and not _is_list(p)
                and len(text) > 40 and p.alignment == AL.JUSTIFY and not pf.left_indent
                and pf.first_line_indent is None):
            pf.first_line_indent = Mm(P["first_line_cm"] * 10)
            log["thut_dau_dong"] += 1
        for run in p.runs:
            if is_title:
                size = P["title"]
            elif upper or (run.text.strip() and _is_upper_line(run.text)):
                size = P["upper"]
            else:
                size = P["body"]
            _font_run(run, size)
            if opts.get("kinh_gui", True) and text.startswith("Kính gửi") and "Kính gửi" in run.text:
                run.bold = True

    # 4. Bảng
    for ti, table in enumerate(doc.tables, 1):
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    log["o_bang"] += 1
                    fix_runs(p, f"Bảng {ti}")
                    pf = p.paragraph_format
                    if _has_image(p):
                        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    else:
                        _line(pf, P["table_line"])
                    pf.space_before, pf.space_after = Pt(P["table_before"]), Pt(P["table_after"])
                    # Lề trong ô chỉ cho chữ căn trái/đều dài hơn vài ký tự (không làm vỡ cột STT, khối Quốc hiệu)
                    if (P["cell_indent_mm"] and len(p.text.strip()) > 3
                            and _eff_align(p) in (None, AL.LEFT, AL.JUSTIFY, AL.DISTRIBUTE)):
                        pf.left_indent = pf.left_indent if (pf.left_indent and pf.left_indent > Mm(2)) \
                            else Mm(P["cell_indent_mm"])
                        pf.right_indent = Mm(P["cell_indent_mm"])
                    for run in p.runs:
                        _font_run(run, P["table"])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), log


# Giữ tương thích tên hàm cũ (dùng từ dòng lệnh)
def standardize_file(in_path, out_path, profile="texo", paper="thuong", opts=None):
    with open(in_path, "rb") as f:
        data, log = standardize(f.read(), profile, paper, opts)
    with open(out_path, "wb") as f:
        f.write(data)
    return log
